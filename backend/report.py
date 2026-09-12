from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from db import conn
from auth import get_current_user, check_role
import io
from openpyxl import Workbook

router = APIRouter()


# ==================================================
# PURCHASE ORDER REPORT
# ==================================================

@router.get("/reports/purchase-orders")
def purchase_order_report(
    page: int = 1,
    limit: int = 50,
    current_user: dict = Depends(get_current_user)
):
    user_role = current_user.get("role")
    user_vendor_id = current_user.get("vendor_id")
    cursor = None

    try:
        conn.rollback()
        cursor = conn.cursor()
        offset = (page - 1) * limit

        if user_role == "Vendor":
            if not user_vendor_id:
                return []
            cursor.execute("""
                SELECT
                    p.id,
                    p.vendor_id,
                    v.vendor_name,
                    p.product_name,
                    p.quantity,
                    p.unit_price,
                    p.total_amount,
                    p.order_date,
                    p.expected_delivery,
                    p.status
                FROM purchase_orders p
                LEFT JOIN vendors v
                    ON p.vendor_id = v.id
                WHERE p.vendor_id = %s
                ORDER BY
                    p.order_date DESC,
                    p.id DESC
                LIMIT %s OFFSET %s
            """, (user_vendor_id, limit, offset))
        else:
            cursor.execute("""
                SELECT
                    p.id,
                    p.vendor_id,
                    v.vendor_name,
                    p.product_name,
                    p.quantity,
                    p.unit_price,
                    p.total_amount,
                    p.order_date,
                    p.expected_delivery,
                    p.status
                FROM purchase_orders p
                LEFT JOIN vendors v
                    ON p.vendor_id = v.id
                ORDER BY
                    p.order_date DESC,
                    p.id DESC
                LIMIT %s OFFSET %s
            """, (limit, offset))

        rows = cursor.fetchall()

        reports = []

        for row in rows:
            reports.append({
                "id": row[0],
                "vendor_id": row[1],
                "vendor_name": row[2] or "Unknown Vendor",
                "product_name": row[3] or "N/A",
                "quantity": int(row[4] or 0),
                "unit_price": float(row[5] or 0),
                "total_amount": float(row[6] or 0),
                "order_date": str(row[7]) if row[7] else "N/A",
                "expected_delivery": str(row[8]) if row[8] else "N/A",
                "status": row[9] or "Unknown"
            })

        return reports

    except Exception as e:
        conn.rollback()
        print("PURCHASE ORDER REPORT ERROR:", repr(e))
        return {
            "error": str(e)
        }
    finally:
        if cursor:
            cursor.close()


# ==================================================
# PURCHASE ORDER REPORT EXCEL EXPORT
# ==================================================

@router.get("/reports/purchase-orders/excel")
def purchase_order_report_excel(current_user: dict = Depends(get_current_user)):
    data = purchase_order_report(page=1, limit=1000, current_user=current_user)
    if isinstance(data, dict) and "error" in data:
        raise HTTPException(status_code=500, detail=data["error"])

    wb = Workbook()
    ws = wb.active
    ws.title = "Purchase Orders"

    # Header columns
    ws.append([
        "Purchase Order ID", "Vendor ID", "Vendor Name", "Product Name", 
        "Quantity", "Unit Price", "Total Amount", "Order Date", 
        "Expected Delivery", "Status"
    ])

    for item in data:
        ws.append([
            item["id"],
            item["vendor_id"],
            item["vendor_name"],
            item["product_name"],
            item["quantity"],
            item["unit_price"],
            item["total_amount"],
            item["order_date"],
            item["expected_delivery"],
            item["status"]
        ])

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    try:
        from audit_logs import log_action
        log_action(
            user_id=current_user.get("id"),
            user_name=current_user.get("name"),
            user_email=current_user.get("email"),
            action="REPORT_GENERATED",
            entity_type="REPORT",
            entity_id="PO_REPORT",
            details="Generated Excel export for Purchase Orders Report"
        )
    except Exception as le:
        print("Audit log report error:", le)

    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=purchase_order_report.xlsx"}
    )


# ==================================================
# VENDOR RELIABILITY REPORT
# ==================================================

@router.get("/reports/vendor-reliability")
def vendor_reliability_report(current_user: dict = Depends(get_current_user)):
    user_role = current_user.get("role")
    user_vendor_id = current_user.get("vendor_id")
    cursor = None

    try:
        conn.rollback()
        cursor = conn.cursor()

        if user_role == "Vendor":
            if not user_vendor_id:
                return []
            cursor.execute("""
                SELECT
                    v.id,
                    v.vendor_name,
                    v.quality_score,
                    v.delivery_rate,
                    v.reliability_score,
                    COUNT(po.id) AS total_orders,
                    COUNT(CASE WHEN LOWER(po.status) IN ('completed', 'delivered') THEN po.id END) AS completed_orders,
                    COUNT(CASE WHEN LOWER(po.status) = 'pending' THEN po.id END) AS pending_orders,
                    COUNT(CASE WHEN LOWER(po.status) = 'delivered' THEN po.id END) AS delivered_orders
                FROM vendors v
                LEFT JOIN purchase_orders po
                    ON v.id = po.vendor_id
                WHERE v.id = %s
                GROUP BY
                    v.id,
                    v.vendor_name,
                    v.quality_score,
                    v.delivery_rate,
                    v.reliability_score
            """, (user_vendor_id,))
        else:
            cursor.execute("""
                SELECT
                    v.id,
                    v.vendor_name,
                    v.quality_score,
                    v.delivery_rate,
                    v.reliability_score,
                    COUNT(po.id) AS total_orders,
                    COUNT(CASE WHEN LOWER(po.status) IN ('completed', 'delivered') THEN po.id END) AS completed_orders,
                    COUNT(CASE WHEN LOWER(po.status) = 'pending' THEN po.id END) AS pending_orders,
                    COUNT(CASE WHEN LOWER(po.status) = 'delivered' THEN po.id END) AS delivered_orders
                FROM vendors v
                LEFT JOIN purchase_orders po
                    ON v.id = po.vendor_id
                GROUP BY
                    v.id,
                    v.vendor_name,
                    v.quality_score,
                    v.delivery_rate,
                    v.reliability_score
                ORDER BY
                    v.reliability_score DESC
            """)

        rows = cursor.fetchall()

        reports = []

        for row in rows:
            reliability_score = float(row[4] or 0)

            if reliability_score >= 80:
                performance = "Excellent"
                risk = "Low Risk"
                recommendation = "Preferred Vendor"
            elif reliability_score >= 60:
                performance = "Good" if reliability_score >= 70 else "Average"
                risk = "Medium Risk"
                recommendation = "Monitor Vendor"
            else:
                performance = "Poor"
                risk = "High Risk"
                recommendation = "Review Vendor"

            reports.append({
                "vendor_id": row[0],
                "vendor_name": row[1],
                "total_orders": int(row[5] or 0),
                "completed_orders": int(row[6] or 0),
                "pending_orders": int(row[7] or 0),
                "delivered_orders": int(row[8] or 0),
                "quality_score": float(row[2] or 0),
                "delivery_rate": float(row[3] or 0),
                "reliability_score": reliability_score,
                "performance": performance,
                "risk": risk,
                "recommendation": recommendation
            })

        return reports

    except Exception as e:
        conn.rollback()
        print("VENDOR RELIABILITY REPORT ERROR:", repr(e))
        return {
            "error": str(e)
        }
    finally:
        if cursor:
            cursor.close()


# ==================================================
# VENDOR RELIABILITY REPORT EXCEL EXPORT
# ==================================================

@router.get("/reports/vendor-reliability/excel")
def vendor_reliability_report_excel(current_user: dict = Depends(get_current_user)):
    data = vendor_reliability_report(current_user)
    if isinstance(data, dict) and "error" in data:
        raise HTTPException(status_code=500, detail=data["error"])

    wb = Workbook()
    ws = wb.active
    ws.title = "Vendor Reliability"

    # Header columns
    ws.append([
        "Vendor ID", "Vendor Name", "Total Orders", "Completed Orders", 
        "Pending Orders", "Delivered Orders", "Quality Score", 
        "Delivery Rate", "Reliability Score", "Performance", 
        "Risk", "Recommendation"
    ])

    for item in data:
        ws.append([
            item["vendor_id"],
            item["vendor_name"],
            item["total_orders"],
            item["completed_orders"],
            item["pending_orders"],
            item["delivered_orders"],
            item["quality_score"],
            item["delivery_rate"],
            item["reliability_score"],
            item["performance"],
            item["risk"],
            item["recommendation"]
        ])

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    try:
        from audit_logs import log_action
        log_action(
            user_id=current_user.get("id"),
            user_name=current_user.get("name"),
            user_email=current_user.get("email"),
            action="REPORT_GENERATED",
            entity_type="REPORT",
            entity_id="RELIABILITY_REPORT",
            details="Generated Excel export for Vendor Reliability Report"
        )
    except Exception as le:
        print("Audit log report error:", le)

    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=vendor_reliability_report.xlsx"}
    )


# ==================================================
# PROCUREMENT REPORT
# ==================================================
@router.get("/reports/procurement")
def procurement_report(
    page: int = 1,
    limit: int = 50,
    current_user: dict = Depends(get_current_user)
):
    user_role = current_user.get("role")
    user_vendor_id = current_user.get("vendor_id")
    offset = (page - 1) * limit
    try:
        conn.rollback()
        with conn.cursor() as cur:
            if user_role == "Vendor":
                if not user_vendor_id:
                    return []
                cur.execute("""
                    SELECT pr.id, pr.vendor_id, COALESCE(v.vendor_name, 'Unassigned') AS vendor_name,
                           pr.product_name, pr.quantity, pr.request_date, pr.requested_by, pr.status, pr.purchase_order_id,
                           COALESCE(pr.unit_price, 0) AS unit_price, COALESCE(pr.total_amount, 0) AS total_amount
                    FROM purchase_requests pr
                    LEFT JOIN vendors v ON pr.vendor_id = v.id
                    WHERE pr.vendor_id = %s
                    ORDER BY pr.id DESC
                    LIMIT %s OFFSET %s
                """, (user_vendor_id, limit, offset))
            else:
                cur.execute("""
                    SELECT pr.id, pr.vendor_id, COALESCE(v.vendor_name, 'Unassigned') AS vendor_name,
                           pr.product_name, pr.quantity, pr.request_date, pr.requested_by, pr.status, pr.purchase_order_id,
                           COALESCE(pr.unit_price, 0) AS unit_price, COALESCE(pr.total_amount, 0) AS total_amount
                    FROM purchase_requests pr
                    LEFT JOIN vendors v ON pr.vendor_id = v.id
                    ORDER BY pr.id DESC
                    LIMIT %s OFFSET %s
                """, (limit, offset))
            rows = cur.fetchall()
        return [{
            "id": r[0],
            "vendor_id": r[1],
            "vendor_name": r[2],
            "product_name": r[3],
            "quantity": int(r[4] or 0),
            "request_date": str(r[5]) if r[5] else "",
            "requested_by": r[6] or "N/A",
            "status": r[7] or "Pending",
            "purchase_order_id": r[8],
            "unit_price": float(r[9] or 0),
            "total_amount": float(r[10] or 0)
        } for r in rows]
    except Exception as e:
        print("PROCUREMENT REPORT ERROR:", e)
        return {"error": str(e)}

@router.get("/reports/procurement/excel")
def procurement_report_excel(current_user: dict = Depends(get_current_user)):
    data = procurement_report(page=1, limit=1000, current_user=current_user)
    if isinstance(data, dict) and "error" in data:
        raise HTTPException(status_code=500, detail=data["error"])
    wb = Workbook()
    ws = wb.active
    ws.title = "Procurement Requisitions"
    ws.append(["Request ID", "Vendor ID", "Vendor Name", "Product Name", "Quantity", "Request Date", "Requested By", "Status", "Linked PO ID"])
    for item in data:
        ws.append([item["id"], item["vendor_id"], item["vendor_name"], item["product_name"], item["quantity"], item["request_date"], item["requested_by"], item["status"], item["purchase_order_id"]])
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=procurement_report.xlsx"}
    )

# ==================================================
# COMPLIANCE REPORT (Quality Inspections)
# ==================================================
@router.get("/reports/compliance")
def compliance_report(
    page: int = 1,
    limit: int = 50,
    current_user: dict = Depends(get_current_user)
):
    user_role = current_user.get("role")
    user_vendor_id = current_user.get("vendor_id")
    offset = (page - 1) * limit
    try:
        conn.rollback()
        with conn.cursor() as cur:
            if user_role == "Vendor":
                if not user_vendor_id:
                    return []
                cur.execute("""
                    SELECT q.id, q.vendor_id, COALESCE(v.vendor_name, 'Unknown') AS vendor_name,
                           q.purchase_order_id, q.inspection_date, q.quantity_inspected, q.quantity_passed,
                           q.quantity_failed, q.quality_score, q.inspection_status, q.remarks
                    FROM quality_inspections q
                    LEFT JOIN vendors v ON q.vendor_id = v.id
                    WHERE q.vendor_id = %s
                    ORDER BY q.id DESC
                    LIMIT %s OFFSET %s
                """, (user_vendor_id, limit, offset))
            else:
                cur.execute("""
                    SELECT q.id, q.vendor_id, COALESCE(v.vendor_name, 'Unknown') AS vendor_name,
                           q.purchase_order_id, q.inspection_date, q.quantity_inspected, q.quantity_passed,
                           q.quantity_failed, q.quality_score, q.inspection_status, q.remarks
                    FROM quality_inspections q
                    LEFT JOIN vendors v ON q.vendor_id = v.id
                    ORDER BY q.id DESC
                    LIMIT %s OFFSET %s
                """, (limit, offset))
            rows = cur.fetchall()
        return [{
            "id": r[0],
            "vendor_id": r[1],
            "vendor_name": r[2],
            "purchase_order_id": r[3],
            "inspection_date": str(r[4]) if r[4] else "",
            "quantity_inspected": int(r[5] or 0),
            "quantity_passed": int(r[6] or 0),
            "quantity_failed": int(r[7] or 0),
            "quality_score": float(r[8] or 0),
            "inspection_status": r[9] or "Pending",
            "remarks": r[10] or ""
        } for r in rows]
    except Exception as e:
        print("COMPLIANCE REPORT ERROR:", e)
        return {"error": str(e)}

@router.get("/reports/compliance/excel")
def compliance_report_excel(current_user: dict = Depends(get_current_user)):
    data = compliance_report(page=1, limit=1000, current_user=current_user)
    if isinstance(data, dict) and "error" in data:
        raise HTTPException(status_code=500, detail=data["error"])
    wb = Workbook()
    ws = wb.active
    ws.title = "Quality Compliance"
    ws.append(["Inspection ID", "Vendor ID", "Vendor Name", "PO ID", "Inspection Date", "Inspected Qty", "Passed Qty", "Failed Qty", "Quality Score", "Status", "Remarks"])
    for item in data:
        ws.append([item["id"], item["vendor_id"], item["vendor_name"], item["purchase_order_id"], item["inspection_date"], item["quantity_inspected"], item["quantity_passed"], item["quantity_failed"], item["quality_score"], item["inspection_status"], item["remarks"]])
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=compliance_report.xlsx"}
    )

# ==================================================
# CONTRACT REPORT
# ==================================================
@router.get("/reports/contracts")
def contracts_report(
    page: int = 1,
    limit: int = 50,
    current_user: dict = Depends(get_current_user)
):
    user_role = current_user.get("role")
    user_vendor_id = current_user.get("vendor_id")
    offset = (page - 1) * limit
    try:
        conn.rollback()
        with conn.cursor() as cur:
            if user_role == "Vendor":
                if not user_vendor_id:
                    return []
                cur.execute("""
                    SELECT c.id, c.vendor_id, COALESCE(v.vendor_name, 'Unknown') AS vendor_name,
                           c.contract_name, c.start_date, c.end_date, c.status, c.contract_value, c.compliance_status
                    FROM contracts c
                    LEFT JOIN vendors v ON c.vendor_id = v.id
                    WHERE c.vendor_id = %s
                    ORDER BY c.id DESC
                    LIMIT %s OFFSET %s
                """, (user_vendor_id, limit, offset))
            else:
                cur.execute("""
                    SELECT c.id, c.vendor_id, COALESCE(v.vendor_name, 'Unknown') AS vendor_name,
                           c.contract_name, c.start_date, c.end_date, c.status, c.contract_value, c.compliance_status
                    FROM contracts c
                    LEFT JOIN vendors v ON c.vendor_id = v.id
                    ORDER BY c.id DESC
                    LIMIT %s OFFSET %s
                """, (limit, offset))
            rows = cur.fetchall()
        return [{
            "id": r[0],
            "vendor_id": r[1],
            "vendor_name": r[2],
            "contract_name": r[3],
            "start_date": str(r[4]) if r[4] else "",
            "end_date": str(r[5]) if r[5] else "",
            "status": r[6] or "Active",
            "contract_value": float(r[7] or 0),
            "compliance_status": r[8] or "Compliant"
        } for r in rows]
    except Exception as e:
        print("CONTRACTS REPORT ERROR:", e)
        return {"error": str(e)}

@router.get("/reports/contracts/excel")
def contracts_report_excel(current_user: dict = Depends(get_current_user)):
    data = contracts_report(page=1, limit=1000, current_user=current_user)
    if isinstance(data, dict) and "error" in data:
        raise HTTPException(status_code=500, detail=data["error"])
    wb = Workbook()
    ws = wb.active
    ws.title = "Contracts Ledger"
    ws.append(["Contract ID", "Vendor ID", "Vendor Name", "Contract Name", "Start Date", "End Date", "Status", "Contract Value", "Compliance"])
    for item in data:
        ws.append([item["id"], item["vendor_id"], item["vendor_name"], item["contract_name"], item["start_date"], item["end_date"], item["status"], item["contract_value"], item["compliance_status"]])
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=contracts_report.xlsx"}
    )

# ==================================================
# FINANCIAL REPORT
# ==================================================
@router.get("/reports/financial")
def financial_report(
    page: int = 1,
    limit: int = 50,
    current_user: dict = Depends(get_current_user)
):
    user_role = current_user.get("role")
    user_vendor_id = current_user.get("vendor_id")
    offset = (page - 1) * limit
    try:
        conn.rollback()
        with conn.cursor() as cur:
            if user_role == "Vendor":
                if not user_vendor_id:
                    return []
                cur.execute("""
                    SELECT i.id, i.invoice_number, i.po_id, i.vendor_id, COALESCE(v.vendor_name, 'Unknown') AS vendor_name,
                           i.invoice_amount, i.invoice_date, i.due_date, i.payment_status, i.payment_date
                    FROM invoices i
                    LEFT JOIN vendors v ON i.vendor_id = v.id
                    WHERE i.vendor_id = %s
                    ORDER BY i.id DESC
                    LIMIT %s OFFSET %s
                """, (user_vendor_id, limit, offset))
            else:
                cur.execute("""
                    SELECT i.id, i.invoice_number, i.po_id, i.vendor_id, COALESCE(v.vendor_name, 'Unknown') AS vendor_name,
                           i.invoice_amount, i.invoice_date, i.due_date, i.payment_status, i.payment_date
                    FROM invoices i
                    LEFT JOIN vendors v ON i.vendor_id = v.id
                    ORDER BY i.id DESC
                    LIMIT %s OFFSET %s
                """, (limit, offset))
            rows = cur.fetchall()
        return [{
            "id": r[0],
            "invoice_number": r[1],
            "po_id": r[2],
            "vendor_id": r[3],
            "vendor_name": r[4],
            "invoice_amount": float(r[5] or 0),
            "invoice_date": str(r[6]) if r[6] else "",
            "due_date": str(r[7]) if r[7] else "",
            "payment_status": r[8] or "Pending",
            "payment_date": str(r[9]) if r[9] else ""
        } for r in rows]
    except Exception as e:
        print("FINANCIAL REPORT ERROR:", e)
        return {"error": str(e)}

@router.get("/reports/financial/excel")
def financial_report_excel(current_user: dict = Depends(get_current_user)):
    data = financial_report(page=1, limit=1000, current_user=current_user)
    if isinstance(data, dict) and "error" in data:
        raise HTTPException(status_code=500, detail=data["error"])
    wb = Workbook()
    ws = wb.active
    ws.title = "Financial Ledger"
    ws.append(["Invoice ID", "Invoice Number", "PO ID", "Vendor ID", "Vendor Name", "Invoice Amount", "Invoice Date", "Due Date", "Payment Status", "Payment Date"])
    for item in data:
        ws.append([item["id"], item["invoice_number"], item["po_id"], item["vendor_id"], item["vendor_name"], item["invoice_amount"], item["invoice_date"], item["due_date"], item["payment_status"], item["payment_date"]])
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=financial_report.xlsx"}
    )

# ==================================================
# AUDIT REPORT
# ==================================================
@router.get("/reports/audit")
def audit_report(
    page: int = 1,
    limit: int = 50,
    current_user: dict = Depends(check_role(["Admin", "Administrator", "Auditor"]))
):
    offset = (page - 1) * limit
    try:
        conn.rollback()
        with conn.cursor() as cur:
            cur.execute("""
                SELECT id, user_id, user_name, user_email, action, entity_type, entity_id, details, ip_address, created_at
                FROM audit_logs
                ORDER BY id DESC
                LIMIT %s OFFSET %s
            """, (limit, offset))
            rows = cur.fetchall()
        return [{
            "id": r[0],
            "user_id": r[1],
            "user_name": r[2] or "System",
            "user_email": r[3] or "N/A",
            "action": r[4],
            "entity_type": r[5],
            "entity_id": r[6],
            "details": r[7],
            "ip_address": r[8] or "127.0.0.1",
            "created_at": str(r[9]) if r[9] else ""
        } for r in rows]
    except Exception as e:
        print("AUDIT REPORT ERROR:", e)
        return {"error": str(e)}

@router.get("/reports/audit/excel")
def audit_report_excel(current_user: dict = Depends(check_role(["Admin", "Administrator", "Auditor"]))):
    data = audit_report(page=1, limit=1000, current_user=current_user)
    if isinstance(data, dict) and "error" in data:
        raise HTTPException(status_code=500, detail=data["error"])
    wb = Workbook()
    ws = wb.active
    ws.title = "Audit Trail Ledger"
    ws.append(["Log ID", "User ID", "User Name", "User Email", "Action", "Entity Type", "Entity ID", "Details", "IP Address", "Timestamp"])
    for item in data:
        ws.append([item["id"], item["user_id"], item["user_name"], item["user_email"], item["action"], item["entity_type"], item["entity_id"], item["details"], item["ip_address"], item["created_at"]])
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=audit_report.xlsx"}
    )

# ==================================================
# AUDIT FINDINGS REPORT
# ==================================================
@router.get("/reports/audit-findings")
def audit_findings_report(
    page: int = 1,
    limit: int = 50,
    current_user: dict = Depends(check_role(["Admin", "Administrator", "Auditor"]))
):
    offset = (page - 1) * limit
    try:
        conn.rollback()
        with conn.cursor() as cur:
            cur.execute("""
                SELECT finding_code, audit_area, vendor_name, entity_type, entity_id, 
                       risk_level, identified_date, audit_status, resolution_status, description
                FROM audit_findings
                ORDER BY 
                    CASE WHEN LOWER(risk_level) = 'critical' THEN 1
                         WHEN LOWER(risk_level) = 'high' THEN 2
                         WHEN LOWER(risk_level) = 'medium' THEN 3
                         ELSE 4 END ASC,
                    identified_date DESC,
                    id DESC
                LIMIT %s OFFSET %s
            """, (limit, offset))
            rows = cur.fetchall()
        return [{
            "finding_code": r[0],
            "audit_area": r[1],
            "vendor_name": r[2],
            "entity_type": r[3],
            "entity_id": r[4],
            "risk_level": r[5],
            "identified_date": str(r[6]) if r[6] else "",
            "audit_status": r[7] or "Open",
            "resolution_status": r[8] or "Unresolved",
            "description": r[9]
        } for r in rows]
    except Exception as e:
        print("AUDIT FINDINGS REPORT ERROR:", e)
        return {"error": str(e)}

@router.get("/reports/audit-findings/excel")
def audit_findings_report_excel(current_user: dict = Depends(check_role(["Admin", "Administrator", "Auditor"]))):
    data = audit_findings_report(page=1, limit=1000, current_user=current_user)
    if isinstance(data, dict) and "error" in data:
        raise HTTPException(status_code=500, detail=data["error"])
    wb = Workbook()
    ws = wb.active
    ws.title = "Audit Findings Ledger"
    ws.append(["Finding ID", "Audit Area", "Vendor Partner", "Entity Type", "Entity ID", "Risk Level", "Identified Date", "Audit Status", "Resolution Status", "Description"])
    for item in data:
        ws.append([
            item["finding_code"],
            item["audit_area"],
            item["vendor_name"],
            item["entity_type"],
            item["entity_id"],
            item["risk_level"],
            item["identified_date"],
            item["audit_status"],
            item["resolution_status"],
            item["description"]
        ])
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    try:
        from audit_logs import log_action
        log_action(
            user_id=current_user.get("id"),
            user_name=current_user.get("name"),
            user_email=current_user.get("email"),
            action="REPORT_GENERATED",
            entity_type="REPORT",
            entity_id="FINDINGS_REPORT",
            details="Generated Excel export for Audit Findings Report"
        )
    except Exception as le:
        print("Audit log report error:", le)

    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=audit_findings_report.xlsx"}
    )